"""
parser.py — Handles all resume file parsing (PDF/DOCX).
Extracts:
- raw text
- hyperlinks
- metadata (pages, type, etc.)

This file contains ZERO LLM code.
It is pure file parsing logic.
"""

import os
import re
import fitz  # PyMuPDF
from typing import Dict, Any, List
from docx import Document
from langchain.tools import tool


# Regex to catch inline hyperlinks inside text
URL_RE = re.compile(r'https?://[^\s\)\]]+')


# ---------------------------------------------------------
# PDF Parsing
# ---------------------------------------------------------

def parse_pdf_blocks(path: str) -> Dict[str, Any]:
    """
    Extract blocks and links from PDF using PyMuPDF.
    Returns: dict with pages, blocks_by_page, page_links
    """
    doc = fitz.open(path)
    pages_blocks, page_links = [], []

    for pno in range(len(doc)):
        page = doc[pno]

        # Text blocks
        blocks = page.get_text("blocks")
        structured = []
        for b in blocks:
            x0, y0, x1, y1, text = b[:5]
            structured.append({
                "bbox": [float(x0), float(y0), float(x1), float(y1)],
                "text": text.strip()
            })

        # Hyperlinks
        links = []
        for link in page.get_links():
            if link.get("uri"):
                rect = link.get("from")
                rect_list = (
                    [float(rect.x0), float(rect.y0), float(rect.x1), float(rect.y1)]
                    if rect else None
                )
                links.append({"uri": link.get("uri"), "from": rect_list})

        pages_blocks.append(structured)
        page_links.append(links)

    return {
        "pages": len(doc),
        "blocks_by_page": pages_blocks,
        "page_links": page_links
    }


def blocks_to_text(blocks_by_page: List[List[Dict[str, Any]]], tol: float = 10.0) -> str:
    """
    Convert PDF block layout into readable text with spacing.
    """
    pages_text = []
    for blocks in blocks_by_page:
        sorted_blocks = sorted(
            blocks,
            key=lambda b: (round(b['bbox'][1], 1), round(b['bbox'][0], 1))
        )

        lines, prev_y = [], None
        for block in sorted_blocks:
            y_top = block['bbox'][1]
            txt = block["text"]

            if not txt:
                continue

            if prev_y is None:
                lines.append(txt)
            else:
                if abs(y_top - prev_y) < tol:
                    # Merge blocks on same line
                    lines[-1] = lines[-1].rstrip() + " " + txt
                else:
                    lines.append(txt)

            prev_y = y_top

        pages_text.append("\n\n".join(lines))

    return "\n\n".join(pages_text)


def parse_pdf_to_text(path: str) -> Dict[str, Any]:
    """
    Turns a PDF into raw_text + metadata (type, pages, hyperlinks)
    """
    parsed = parse_pdf_blocks(path)
    raw_text = blocks_to_text(parsed["blocks_by_page"])

    # Extract hyperlinks
    links = []
    for pno, page_links in enumerate(parsed.get("page_links", [])):
        for l in page_links:
            links.append({
                "uri": l["uri"],
                "page": pno,
                "from": l["from"]
            })

    # Inline URL detection from text
    for m in URL_RE.findall(raw_text):
        if not any(m == ex["uri"] for ex in links):
            links.append({"uri": m, "page": None, "from": None})

    metadata = {
        "type": "pdf",
        "pages": parsed["pages"],
        "hyperlinks": links
    }

    return {"raw_text": raw_text, "metadata": metadata}


# ---------------------------------------------------------
# DOCX Parsing
# ---------------------------------------------------------

def parse_docx_to_text(path: str) -> Dict[str, Any]:
    """
    PRO DOCX PARSER
    Extract:
      ✓ Paragraphs
      ✓ Tables
      ✓ Multi-run text
      ✓ Hyperlinks
      ✓ Headers & footers
    """
    doc = Document(path)
    all_text_parts = []

    # --------- PARAGRAPHS (runs) ----------
    for p in doc.paragraphs:
        if p.runs:
            txt = "".join(run.text for run in p.runs).strip()
        else:
            txt = p.text.strip()

        if txt:
            all_text_parts.append(txt)

    # --------- TABLES ----------
    for table in doc.tables:
        for row in table.rows:
            row_cells = []
            for cell in row.cells:
                cell_text = []
                for p in cell.paragraphs:
                    if p.runs:
                        one = "".join(run.text for run in p.runs).strip()
                    else:
                        one = p.text.strip()
                    if one:
                        cell_text.append(one)

                if cell_text:
                    row_cells.append(" ".join(cell_text))

            if row_cells:
                all_text_parts.append(" | ".join(row_cells))

    # --------- HEADERS ----------
    try:
        hdr = doc.sections[0].header
        for p in hdr.paragraphs:
            t = p.text.strip()
            if t:
                all_text_parts.append(t)
    except:
        pass

    # --------- FOOTERS ----------
    try:
        ftr = doc.sections[0].footer
        for p in ftr.paragraphs:
            t = p.text.strip()
            if t:
                all_text_parts.append(t)
    except:
        pass

    raw_text = "\n".join(all_text_parts)

    # --------- HYPERLINKS ----------
    links = []
    try:
        for rel in doc.part._rels.values():
            if "hyperlink" in rel.reltype:
                links.append({"uri": rel.target_ref, "page": None, "from": None})
    except:
        pass

    # Inline URLs
    for m in URL_RE.findall(raw_text):
        if not any(m == ex["uri"] for ex in links):
            links.append({"uri": m, "page": None, "from": None})

    return {
        "raw_text": raw_text,
        "metadata": {
            "type": "docx",
            "hyperlinks": links
        }
    }


# ---------------------------------------------------------
# PUBLIC FUNCTION
# ---------------------------------------------------------

def parse_resume_file(path: str) -> Dict[str, Any]:
    """
    Main entry point for resume parsing.
    Supports: PDF, DOCX
    """
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        return parse_pdf_to_text(path)

    if ext == ".docx":
        return parse_docx_to_text(path)

    raise ValueError("Unsupported extension: " + ext)



@tool
def parse_resume_tool(file_path: str) -> Dict[str, Any]:
    """
    LangChain tool wrapper for parse_resume_file.
    Input: file_path (str)
    Output: { raw_text, metadata, hyperlinks }
    """
    result = parse_resume_file(file_path)

    return {
        "raw_text": result.get("raw_text", ""),
        "hyperlinks": result.get("metadata", {}).get("hyperlinks", []),
        "metadata": result.get("metadata", {})
    }

# ---------------------------------------------------------------------------------------------------------------------------------------------------
# SELF-TEST (only runs when executing parser.py directly)
# -------------------------------------------------------------------------------------------------------------------------------------------------

if __name__ == "__main__":
    print("\n[parser.py] Self-test started...\n")

    test_path = "D:\\internship\\Viveha-Ai\\agents\\Resume-Agent\\shub_resume.pdf"  # change to an available file path

    if not os.path.exists(test_path):
        print(f"Test file not found: {test_path}")
    else:
        out = parse_resume_file(test_path)
        print("Parsed resume metadata:")
        print(out["metadata"])
        print("\nRaw text preview:")
        print(out["raw_text"][:])

    print("\n[parser.py] Self-test finished.\n")

